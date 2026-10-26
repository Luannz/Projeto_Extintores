import io
import re
import zipfile

import pandas as pd
from docx import Document

regex_chave = re.compile(r"{{([^{}]+)}}")

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def _texto(valor):
    """Converte a célula em texto; célula vazia vira '' (e não 'nan')."""
    if valor is None or pd.isna(valor):
        return ""
    return str(valor).strip()


def formatar_data(valor):
    try:
        data = pd.to_datetime(valor, dayfirst=True, errors="coerce")
    except (ValueError, TypeError):
        return ""
    return "" if pd.isna(data) else data.strftime("%d/%m/%Y")


def formatar_data_extenso(valor):
    """Ex.: 15 de março de 2025 (sem depender do locale do servidor)."""
    try:
        data = pd.to_datetime(valor, dayfirst=True, errors="coerce")
    except (ValueError, TypeError):
        return ""
    if pd.isna(data):
        return ""
    return f"{data.day:02d} de {MESES[data.month - 1]} de {data.year}"


def substituir_em_paragrafos(paragrafos, substituicoes):
    for paragrafo in paragrafos:
        for run in paragrafo.runs:
            if regex_chave.search(run.text):
                texto = run.text
                for chave, valor in substituicoes.items():
                    if chave == "DESCRICAO":
                        texto = texto.replace(f"{{{{{chave}}}}}", valor)
                    else:
                        texto = texto.replace(f"{{{{{chave}}}}}", valor.upper())
                run.text = texto
                run.bold = True


def substituir_campos(doc, substituicoes):
    substituir_em_paragrafos(doc.paragraphs, substituicoes)
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                substituir_em_paragrafos(celula.paragraphs, substituicoes)


def gerar_zip_documentos(arquivo_planilha, arquivo_modelo):
    """Recebe os dois uploads e devolve um ZIP (em memória) com um .docx por linha."""
    df = pd.read_excel(arquivo_planilha)
    modelo_bytes = arquivo_modelo.read()

    buffer = io.BytesIO()
    nomes_usados = {}

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
        for i, row in df.iterrows():
            substituicoes = {
                "FUNCIONARIO": _texto(row.get("FUNCIONARIO")),
                "FUNCAO": _texto(row.get("FUNCAO")),
                "SETOR": _texto(row.get("SETOR")),
                "REALIZA": formatar_data(row.get("REALIZA")),
                "EMPRESA": _texto(row.get("EMPRESA")),
                "DATA": formatar_data_extenso(row.get("DATA")),
                "CPF": _texto(row.get("CPF")),
                "DESCRICAO": _texto(row.get("DESCRICAO")),
            }

            doc = Document(io.BytesIO(modelo_bytes))
            substituir_campos(doc, substituicoes)

            saida = io.BytesIO()
            doc.save(saida)

            nome = substituicoes["FUNCIONARIO"] or f"documento_{i}"
            nome = re.sub(r"[^\w\-]", "_", nome)  # tira espaços e caracteres inválidos

            # dois funcionários com o mesmo nome não se sobrescrevem
            nomes_usados[nome] = nomes_usados.get(nome, 0) + 1
            if nomes_usados[nome] > 1:
                nome = f"{nome}_{nomes_usados[nome]}"

            zipf.writestr(f"{nome}.docx", saida.getvalue())

    buffer.seek(0)
    return buffer